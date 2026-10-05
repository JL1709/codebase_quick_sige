"""Build the prototype DOCX library from retained SiGeKo references and the approved design.

Run with the bundled Python runtime and pass the original reference directory.
Outputs are static assets; document generation in the application remains unchanged.
"""

import argparse
import json
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "src/data/wordTemplateCatalog.json"
APPROVED = ROOT / "artifacts/word-template-design/Brandschutzordnung-Teil-A-Vorlage.docx"
OUTPUT = ROOT / "public/word-templates"
CONTENT_WIDTH_MM = 186
BODY_FONT_PT = 10.5
COMPACT_FORM_FONT_PT = 9.5
HEADING_FONT_PT = 12
TITLE_FONT_PT = 20
FOOTER_MARGIN_MM = 35
CELL_PADDING_TWIPS = 75
COMPACT_CELL_VERTICAL_PADDING_TWIPS = 45
GRID_COLOR = "D2D8DC"
SIGNATURE_ROW_HEIGHT_MM = 20


def run_font(run, size, bold=None, color="000000"):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.highlight_color = None
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attribute), "Arial")


def paragraph_text(paragraph, value, size=BODY_FONT_PT, bold=False, after=3):
    paragraph.text = value
    paragraph.style = "Normal"
    formatting = paragraph.paragraph_format
    formatting.space_before = Pt(0)
    formatting.space_after = Pt(after)
    formatting.line_spacing = 1.1
    formatting.keep_with_next = False
    formatting.widow_control = True
    for run in paragraph.runs:
        run_font(run, size, bold)
    return paragraph


def title(document, value):
    paragraph = paragraph_text(document.add_paragraph(), value, TITLE_FONT_PT, True, 8)
    paragraph.style = "Title"
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def heading(document, value):
    paragraph = paragraph_text(document.add_paragraph(), value, HEADING_FONT_PT, True, 4)
    paragraph.style = "Heading 1"
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def minimal_paragraph(document):
    paragraph = paragraph_text(document.add_paragraph(), "", 1, after=0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    return paragraph


def set_cell_padding(cell):
    properties = cell._tc.get_or_add_tcPr()
    for tag in ("tcMar", "tcBorders"):
        previous = properties.find(qn("w:" + tag))
        if previous is not None:
            properties.remove(previous)
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), str(CELL_PADDING_TWIPS))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def format_table(table, widths=None):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    for tag in ("tblBorders", "tblInd", "tblpPr", "tblCellSpacing"):
        previous = properties.find(qn("w:" + tag))
        if previous is not None:
            properties.remove(previous)
    table_width = properties.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(Mm(CONTENT_WIDTH_MM).twips))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": GRID_COLOR}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    properties.append(borders)
    grid = table._tbl.tblGrid
    old_widths = [int(column.get(qn("w:w"))) for column in grid]
    desired = [Mm(width).twips for width in widths] if widths else [round(width * Mm(CONTENT_WIDTH_MM).twips / sum(old_widths)) for width in old_widths]
    for column, width in zip(grid, desired):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        for node in list(row_properties):
            if node.tag in (qn("w:trHeight"), qn("w:cantSplit")):
                row_properties.remove(node)
        row_properties.append(OxmlElement("w:cantSplit"))
        grid_index = 0
        for cell_element in row._tr.findall(qn("w:tc")):
            cell_properties = cell_element.get_or_add_tcPr()
            span = cell_properties.find(qn("w:gridSpan"))
            count = int(span.get(qn("w:val"))) if span is not None else 1
            width = sum(desired[grid_index:grid_index + count])
            cell_properties.get_or_add_tcW().set(qn("w:w"), str(width))
            grid_index += count
        for cell in row.cells:
            set_cell_padding(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_table(document, rows, widths, font_size=BODY_FONT_PT, header=False):
    table = document.add_table(rows=len(rows), cols=len(widths))
    format_table(table, widths)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            paragraph_text(table.cell(row_index, column_index).paragraphs[0], value, font_size, header and row_index == 0, 1)
        if header and row_index == 0:
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell in table.rows[0].cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F4F6F7")
                cell._tc.get_or_add_tcPr().append(shading)
    return table


def initialize_document(path, approved):
    document = Document(path)
    for relationship in list(document.part.rels.values()):
        if relationship.reltype in (RELATIONSHIP_TYPE.HEADER, RELATIONSHIP_TYPE.FOOTER):
            document.part.drop_rel(relationship.rId)
    for reference_tag in ("headerReference", "footerReference"):
        for reference in document.element.body.findall(".//" + qn("w:" + reference_tag)):
            reference.getparent().remove(reference)
    # Retained section breaks carry old margins and footer references into copied content.
    for section_properties in document.element.body.findall(".//" + qn("w:pPr") + "/" + qn("w:sectPr")):
        section_properties.getparent().remove(section_properties)
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(12)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(FOOTER_MARGIN_MM)
    section.header_distance, section.footer_distance = Mm(9), Mm(7)
    section.different_first_page_header_footer = False
    for area in (section.header, section.footer):
        for element in list(area._element):
            area._element.remove(element)
        area.add_paragraph()
    for style in approved.styles:
        if style.name in ("Company Footer", "Inline Pictogram", "Organisation Logo") and style.name not in document.styles:
            document.styles._element.append(deepcopy(style._element))
    for name, size in (("Normal", BODY_FONT_PT), ("Title", TITLE_FONT_PT), ("Heading 1", HEADING_FONT_PT), ("Heading 2", HEADING_FONT_PT), ("Caption", 8.5), ("Footer", 7.5)):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = document.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        properties = style._element.find(qn("w:pPr"))
        if properties is not None:
            for node in list(properties):
                if node.tag in (qn("w:pBdr"), qn("w:pageBreakBefore")):
                    properties.remove(node)
    return document


def add_brand_and_project(document, approved):
    for index, table in enumerate(approved.tables[:2]):
        document.element.body.insert(-1, deepcopy(table._tbl))
        paragraph = minimal_paragraph(document)
        paragraph.paragraph_format.line_spacing = Pt(5 if index == 0 else 8)


def add_footer(document, approved, name):
    footer = document.sections[0].footer
    for child in list(footer._element):
        footer._element.remove(child)
    for child in approved.sections[0].footer._element:
        footer._element.append(deepcopy(child))
    details = footer.tables[0].cell(0, 2)
    for paragraph in details.paragraphs:
        if paragraph.text in ("Brandschutzordnung Teil A", "DIN 14096"):
            if paragraph.text.startswith("Brandschutzordnung"):
                for run in paragraph.runs:
                    if run.text:
                        run.text = name
                        break
            else:
                paragraph._p.getparent().remove(paragraph._p)


def normalize_paragraph(paragraph, inside_table=False):
    visible_runs = [run for run in paragraph.runs if run.text.strip()]
    is_heading = bool(visible_runs) and len(paragraph.text) < 150 and all(run.bold or run.underline for run in visible_runs)
    formatting = paragraph.paragraph_format
    formatting.space_before = Pt(5 if is_heading else 0)
    formatting.space_after = Pt(3 if inside_table else 4)
    formatting.line_spacing = 1.1
    formatting.keep_with_next = is_heading
    formatting.keep_together = len(paragraph.text) < 500
    formatting.widow_control = True
    formatting.page_break_before = False
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    properties = paragraph._p.get_or_add_pPr()
    for node in list(properties):
        if node.tag in (qn("w:pBdr"), qn("w:tabs")):
            properties.remove(node)
    for run in paragraph.runs:
        run_font(run, HEADING_FONT_PT if is_heading else BODY_FONT_PT)
        if is_heading:
            run.font.bold, run.font.underline = True, False
        for break_node in list(run._r.findall(qn("w:br"))):
            if break_node.get(qn("w:type")) == "page":
                run._r.remove(break_node)
    if not paragraph.text.strip() and not paragraph._p.findall(".//" + qn("w:drawing")) and not paragraph._p.findall(".//" + qn("w:pict")):
        formatting.space_before = formatting.space_after = Pt(0)
        formatting.line_spacing = Pt(1)
        for run in paragraph.runs:
            run_font(run, 1)
    if paragraph.text.strip() == "(Stempel)":
        formatting.space_after = Pt(28)


def keep_form_rows_together(table):
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    for row_index, row in enumerate(table.rows):
        labels = " ".join(cell.text for cell in row.cells)
        if "Ort, Datum" in labels and "Stempel, Unterschrift" in labels:
            row.height = Mm(SIGNATURE_ROW_HEIGHT_MM)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if row_index + 1 >= len(table.rows):
            continue
        next_cells = table.rows[row_index + 1]._tr.findall(qn("w:tc"))
        first_properties = next_cells[0].find(qn("w:tcPr")) if next_cells else None
        merge = first_properties.find(qn("w:vMerge")) if first_properties is not None else None
        if merge is not None and merge.get(qn("w:val")) != "restart":
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def normalize_images(root, maximum_height_mm=48):
    for drawing in root.findall(".//" + qn("w:drawing")):
        extent = drawing.find(".//" + qn("wp:extent"))
        if extent is None:
            continue
        width, height = int(extent.get("cx")), int(extent.get("cy"))
        scale = min(1, int(Mm(maximum_height_mm)) / height)
        for node in drawing.findall(".//" + qn("wp:extent")) + drawing.findall(".//" + qn("a:ext")):
            node.set("cx", str(round(width * scale)))
            node.set("cy", str(round(height * scale)))


def remove_legacy_fields(root):
    for paragraph in root.findall(".//" + qn("w:p")):
        children = list(paragraph)
        index = 0
        while index < len(children):
            child = children[index]
            begin = child.find(qn("w:fldChar")) if child.tag == qn("w:r") else None
            if begin is None or begin.get(qn("w:fldCharType")) != "begin":
                index += 1
                continue
            end_index = index + 1
            while end_index < len(children):
                end = children[end_index].find(qn("w:fldChar"))
                if end is not None and end.get(qn("w:fldCharType")) == "end":
                    break
                end_index += 1
            nodes = children[index:end_index + 1]
            instruction = "".join(text.text or "" for node in nodes for text in node.findall(".//" + qn("w:instrText")))
            if instruction.strip().startswith(("DOCVARIABLE", "DOCPROPERTY")):
                replacements = {"KO_PLAN_FAX": "{{qs.organization.fax}}", "KO_PLAN_MAIL": "{{qs.organization.email}}"}
                replacement = next((value for key, value in replacements.items() if key in instruction), "")
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.text = replacement
                run.append(text)
                paragraph.insert(index, run)
                for node in nodes:
                    paragraph.remove(node)
            index = end_index + 1


def emit_source_content(document, source_tables):
    for source_table in source_tables:
        current_rows = []

        def flush_rows():
            if not current_rows:
                return
            element = deepcopy(source_table._tbl)
            for row in list(element.findall(qn("w:tr"))):
                element.remove(row)
            for row in current_rows:
                element.append(deepcopy(row))
            document.element.body.insert(-1, element)
            table = Table(element, document._body)
            format_table(table)
            for paragraph in element.findall(".//" + qn("w:p")):
                normalize_paragraph(Paragraph(paragraph, document._body), inside_table=True)
            keep_form_rows_together(table)
            normalize_images(element)
            current_rows.clear()
            minimal_paragraph(document)

        for row in source_table._tbl.findall(qn("w:tr")):
            cells = row.findall(qn("w:tc"))
            if len(cells) == 1:
                flush_rows()
                for node in cells[0]:
                    if node.tag == qn("w:p"):
                        if not "".join(node.itertext()).strip() and not node.findall(".//" + qn("w:drawing")) and not node.findall(".//" + qn("w:pict")):
                            continue
                        element = deepcopy(node)
                        document.element.body.insert(-1, element)
                        normalize_paragraph(Paragraph(element, document._body))
                        normalize_images(element)
                    elif node.tag == qn("w:tbl"):
                        nested = deepcopy(node)
                        document.element.body.insert(-1, nested)
                        format_table(Table(nested, document._body))
                        for paragraph in nested.findall(".//" + qn("w:p")):
                            normalize_paragraph(Paragraph(paragraph, document._body), inside_table=True)
            else:
                current_rows.append(row)
        flush_rows()


def loop_table(document, path, columns, widths):
    rows = [labels for labels, _ in columns]
    values = ["{{" + path + "." + key + "}}" for _, key in columns]
    blank = [""] * len(columns)
    table = make_table(document, [rows, blank, values, blank], widths, 11, header=True)
    first = table.cell(1, 0)
    begin = first.paragraphs[0]
    begin.text = "{{#" + path + "}}"
    end = table.cell(3, 0).paragraphs[0]
    end.text = "{{/" + path + "}}"
    for row in (table.rows[1], table.rows[3]):
        for cell in row.cells:
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run_font(run, 1)
    return table


def alarm_plan(document):
    heading(document, "Alarmierung und Notfallkontakte")
    loop_table(document, "qs.project.notfallkontakte.kontakte", [("Alarmierung / Funktion", "bezeichnung"), ("Name", "ansprechpartner"), ("Telefon", "telefon")], [68, 65, 53])
    heading(document, "Projektbeteiligte")
    loop_table(document, "qs.project.projektbeteiligte", [("Funktion", "role"), ("Name", "name"), ("Unternehmen", "company"), ("Telefon", "phone")], [40, 48, 53, 45])
    heading(document, "Wichtige Rufnummern")
    make_table(document, [["Funktion", "Name", "Telefon"], *[[label, "", ""] for label in ["Versorger", "Wasser (bei Störung)", "Stromhavarie", "Gewerbeaufsichtsamt", "Berufsgenossenschaft Bauwirtschaft"]]], [68, 65, 53], 10.5, header=True)


def project_site_plan(document):
    for value in ("{{#qs.project.files}}", "{{qs.project.files.image}}", "{{/qs.project.files}}"):
        paragraph = paragraph_text(document.add_paragraph(), value, 1, after=2)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_text(document.add_paragraph(), "Baustellenzufahrt: {{qs.project.allgemein.baustellenzufahrt}}", 11)


def prepare_escape_plan(source_table):
    for drawing_tag in ("w:drawing", "w:pict"):
        for drawing in list(source_table._tbl.findall(".//" + qn(drawing_tag))):
            if any("Beispiel" in (text.text or "") for text in drawing.findall(".//" + qn("w:t"))):
                drawing.getparent().remove(drawing)


def finish_escape_plan(document):
    paragraphs = [Paragraph(node, document._body) for node in document.element.body.findall(qn("w:p"))]
    picture = next((paragraph for paragraph in paragraphs if paragraph._p.findall(".//" + qn("w:drawing"))), None)
    if picture is None:
        raise ValueError("The escape plan reference picture is missing")
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drawing = picture._p.find(".//" + qn("w:drawing"))
    extent = drawing.find(".//" + qn("wp:extent"))
    scale = int(Mm(80)) / int(extent.get("cy"))
    for node in drawing.findall(".//" + qn("wp:extent")) + drawing.findall(".//" + qn("a:ext")):
        node.set("cx", str(round(int(node.get("cx")) * scale)))
        node.set("cy", str(round(int(node.get("cy")) * scale)))
    caption = paragraph_text(document.add_paragraph(), "Beispielplan – vor Verwendung projektbezogen ersetzen", 8.5)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture._p.addnext(caption._p)


def fill_known_project_slots(document, slug):
    for node in document.element.body.findall(".//" + qn("w:p")):
        paragraph = Paragraph(node, document._body)
        value = paragraph.text
        replacement = None
        if slug == "baustellenordnung" and value.startswith("Die Baustelle liegt"):
            replacement = "Die Baustelle liegt in {{qs.project.allgemein.projektadresse.strasse}}, {{qs.project.allgemein.projektadresse.ort}}."
        elif slug == "sigeplan-allgemeine-massnahmen" and value.startswith("Das Bauvorhaben umfasst die"):
            replacement = "{{qs.project.allgemein.kurzbeschreibung}}"
        elif slug == "brandschutzordnung-teil-b" and value.startswith("Sie gilt sachlich und räumlich"):
            replacement = value.replace("….", "{{qs.project.allgemein.projektadresse.strasse}}, {{qs.project.allgemein.projektadresse.ort}}.")
        elif slug == "brandschutzordnung-teil-b" and value.startswith("Die Brandschutzordnung Teil B für das Bauvorhaben"):
            replacement = value.replace("…", "{{qs.project.name}}")
        if replacement:
            paragraph_text(paragraph, replacement)


def demolition_form(document, source_table):
    rows = []
    for row in source_table.rows:
        cells = []
        seen = set()
        for cell in row.cells:
            if cell._tc not in seen:
                seen.add(cell._tc)
                cells.append(cell)
        label = cells[0].text.strip()
        if label == "Beginn:":
            rows.append(["Beginn / Ende", "{{qs.project.allgemein.geplanter_beginn}} / {{qs.project.allgemein.geplantes_ende}}"])
        elif label in ("Aufsichtführender:", "LBO-Bauleiter:", "Zuständige BG:"):
            rows.append([label, ""])
            rows.append([cells[2].text.strip(), ""])
        elif label == "Objektart:":
            rows.append([label, "□ Gebäude   □ Brücke   □ Maschine/Anlage   □ Sonstiges"])
        elif any(cell.text.strip() == "nein" for cell in cells):
            extra = " ".join(cell.text.strip() for cell in cells[1:] if cell.text.strip() not in ("", "nein", "ja"))
            rows.append([label, "□ Nein   □ Ja\n" + extra])
        else:
            defaults = {
                "Abbruchstelle (Ort, Straße)": "{{qs.project.allgemein.projektadresse.strasse}}, {{qs.project.allgemein.projektadresse.ort}}",
                "Abbruchgenehmigung, Nr.": "{{qs.project.allgemein.baugenehmigungsnummer}}",
                "Auftraggeber:": "{{qs.project.allgemein.bauherr}}",
                "Kurzbeschreibung der baulichen Anlage:": "{{qs.project.allgemein.kurzbeschreibung}}",
            }
            rows.append([label, defaults.get(label, " ".join(cell.text.strip() for cell in cells[1:] if cell.text.strip()))])
    make_table(document, rows, [66, 120], 10)


def advance_notice(document):
    paragraph_text(document.add_paragraph(), "gemäß § 2 der Verordnung über Sicherheit und Gesundheitsschutz auf Baustellen (Baustellenverordnung – BaustellV)", 10, after=6)
    prefix = "qs.project.vorankuendigung"
    slot = lambda path: "{{" + prefix + "." + path + "}}"
    rows = [
        ["An (zuständiges Gewerbeaufsichtsamt)", slot("behoerde.name") + "\n" + slot("behoerde.adresse") + "\nFax: " + slot("behoerde.fax")],
        ["1. Bezeichnung und Ort der Baustelle", "{{qs.project.name}}\n{{qs.project.allgemein.projektadresse.strasse}}, {{qs.project.allgemein.projektadresse.ort}}"],
        ["2. Name und Anschrift des Bauherren", "{{qs.project.allgemein.bauherr}}\n" + slot("bauherr_adresse")],
        ["3. Anstelle des Bauherren verantwortlicher Dritter", slot("verantwortlicher_dritter.name") + "\n" + slot("verantwortlicher_dritter.adresse")],
        ["4. Art des Bauvorhabens", "{{qs.project.allgemein.kurzbeschreibung}}"],
    ]
    for key, label in [("koordination_planung", "5. Koordination für die Planung der Ausführung"), ("koordination_ausfuehrung", "Koordination für die Ausführung des Bauvorhabens")]:
        value = slot(key + ".name") + "\n" + slot(key + ".adresse") + "\nTel.: " + slot(key + ".telefon") + " · Fax: " + slot(key + ".fax") + "\nE-Mail: " + slot(key + ".e_mail")
        rows.append([label + " (sofern erforderlich)", value])
    rows.extend([
        ["6. Voraussichtlicher Beginn und Ende der Arbeiten", "von {{qs.project.allgemein.geplanter_beginn}} bis {{qs.project.allgemein.geplantes_ende}}"],
        ["7. Höchstzahl gleichzeitig Beschäftigte", slot("hoechstzahl_beschaeftigte")],
        ["8. Voraussichtliche Zahl der Arbeitgeber", slot("zahl_arbeitgeber")],
        ["9. Zahl der Unternehmer ohne Beschäftigte", slot("zahl_unternehmer_ohne_beschaeftigte")],
        ["10. Bereits ausgewählte Arbeitgeber und Unternehmer ohne Beschäftigte", slot("ausgewaehlte_unternehmen") + "\n(weitere Angaben, ggf. als Anlage)"],
    ])
    make_table(document, rows, [68, 118], COMPACT_FORM_FONT_PT)
    paragraph_text(document.add_paragraph(), "Ort/Datum: ____________________    Name: ____________________    Unterschrift: ____________________", 9, after=4)
    paragraph_text(document.add_paragraph(), "(Bauherr oder anstelle des Bauherren verantwortlicher Dritter)", 9)
    paragraph_text(document.add_paragraph(), "Verteiler: 1 x zuständige Behörde · 1 x Baustellenaushang · 1 x Bauherr", 9)


def compact_emergency_checklist(document):
    for table in document.tables[2:]:
        for margin in table._tbl.findall(".//" + qn("w:tcMar")):
            for edge in ("top", "bottom"):
                margin.find(qn("w:" + edge)).set(qn("w:w"), str(COMPACT_CELL_VERTICAL_PADDING_TWIPS))
        for element in table._tbl.findall(".//" + qn("w:p")):
            paragraph = Paragraph(element, document._body)
            formatting = paragraph.paragraph_format
            formatting.space_before = Pt(0)
            formatting.space_after = Pt(1)
            formatting.line_spacing = 1
            for run in paragraph.runs:
                run_font(run, COMPACT_FORM_FONT_PT)


def clean_settings(document, name):
    settings = document.settings._element
    for node in list(settings):
        if node.tag in (qn("w:docVars"), qn("w:rsids"), qn("w:evenAndOddHeaders")):
            settings.remove(node)
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    properties = document.core_properties
    properties.title, properties.author = name, "QuickReports"
    properties.last_modified_by = "QuickReports"
    properties.subject = "Word-Vorlage mit Projektplatzhaltern"
    properties.comments = properties.keywords = ""


def build_entry(entry, reference_directory, approved):
    source_path = reference_directory / entry["sourceFilename"]
    if sha256(source_path.read_bytes()).hexdigest() != entry["sourceSha256"]:
        raise ValueError("Reference changed: " + str(source_path))
    output_path = OUTPUT / entry["filename"]
    if entry["slug"] == "brandschutzordnung-teil-a":
        copyfile(APPROVED, output_path)
        return
    document = initialize_document(source_path, approved)
    tables = list(document.tables)
    content_tables = tables if entry["slug"] == "vorankuendigung" else tables[1:]
    content_tables[0]._tbl.remove(content_tables[0].rows[0]._tr)
    for element in list(document.element.body):
        if element.tag != qn("w:sectPr"):
            document.element.body.remove(element)
    add_brand_and_project(document, approved)
    display_title = {
        "abbruchanweisung": "Muster einer Abbruchanweisung",
        "alarmplan": "Alarmplan und Telefonliste",
        "anschlussbewehrung": "Gefährdungen durch Bewehrung",
        "firmenangaben": "Angaben zu Sicherheit und Gesundheitsschutz",
        "hubarbeitsbuehnen": "Verwendung von Hubarbeitsbühnen",
        "schutznetze": "Verwendung von Schutznetzen",
        "sigeplan-bedeutung": "Sicherheits- und Gesundheitsschutzplan",
        "verwendung-von-psaga": "Persönliche Schutzausrüstung gegen Absturz PSAgA",
    }.get(entry["slug"], entry["name"])
    title(document, display_title)
    if entry["slug"] == "alarmplan":
        alarm_plan(document)
    elif entry["slug"] == "abbruchanweisung":
        demolition_form(document, tables[1])
    elif entry["slug"] == "vorankuendigung":
        advance_notice(document)
    elif entry["slug"] == "lageplan":
        project_site_plan(document)
    else:
        if entry["slug"] == "flucht-rettungsplan":
            prepare_escape_plan(content_tables[0])
        emit_source_content(document, content_tables)
        remove_legacy_fields(document.element.body)
        fill_known_project_slots(document, entry["slug"])
        if entry["slug"] == "flucht-rettungsplan":
            finish_escape_plan(document)
        elif entry["slug"] == "notfallplanung":
            compact_emergency_checklist(document)
        elif entry["slug"] == "sigeplan-bedeutung":
            for paragraph in document.paragraphs:
                if paragraph.text.startswith("Health and safety plan"):
                    paragraph.paragraph_format.page_break_before = True
    minimal_paragraph(document)
    add_footer(document, approved, entry["name"])
    clean_settings(document, entry["name"])
    document.save(output_path)
    assert sha256(source_path.read_bytes()).hexdigest() == entry["sourceSha256"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("reference_directory", type=Path)
    arguments = parser.parse_args()
    OUTPUT.mkdir(exist_ok=True)
    approved = Document(APPROVED)
    for entry in json.loads(CATALOG.read_text()):
        build_entry(entry, arguments.reference_directory, approved)
        print(entry["filename"], flush=True)


if __name__ == "__main__":
    main()
